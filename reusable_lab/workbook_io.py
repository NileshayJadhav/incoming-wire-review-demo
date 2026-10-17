"""Read-only discovery of incoming-wire workbooks. No workbook authoring here."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

INPUT_FIELDS = ("sample_id", "date_time", "customer", "account", "amount")
RESULT_FIELDS = ("fee", "fee_charged", "classification", "travel", "ofac", "posting", "comments")
ALL_FIELDS = INPUT_FIELDS + RESULT_FIELDS
ALIASES = {
    "sample_id": ("#", "sample", "sample id", "sample number", "sample no", "sample #", "item", "item no"),
    "date_time": ("date", "date/time", "date time", "wire date", "transaction date", "transaction date/time"),
    "customer": ("customer", "customer name", "member", "member name", "beneficiary", "beneficiary name"),
    "account": ("account", "account number", "account no", "account #", "acct", "acct number", "member account"),
    "amount": ("amount", "wire amount", "transaction amount", "principal", "principal amount", "amount usd"),
    "fee": ("fee", "wire fee", "fee amount", "recorded fee", "fee usd"),
    "fee_charged": ("fee charged", "fee charged?", "fee charged correctly", "fee result", "fee test"),
    "classification": ("domestic/foreign", "domestic foreign", "classification", "wire classification", "wire type"),
    "travel": ("travel", "travel rule", "travel rule review", "travel result", "travel test"),
    "ofac": ("ofac", "ofac review", "ofac result", "ofac test"),
    "posting": ("transaction posted correctly to acct", "transaction posted correctly to account", "posted correctly", "posting", "posting result", "posting test"),
    "comments": ("comments", "comment", "review comments", "review notes", "notes"),
}


class WorkbookLayoutError(ValueError):
    """Bounded workbook-layout/input error; contains no source row values."""


def require(ok, message):
    if not ok:
        raise WorkbookLayoutError(message)


def source_hash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().casefold()
    if text == "#":
        return "#"
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_id(value):
    require(value is not None and not isinstance(value, bool), "Missing or invalid sample ID.")
    if isinstance(value, (int, float, Decimal)):
        d = Decimal(str(value))
        require(d.is_finite(), "Invalid sample ID.")
        value = format(d, "f").rstrip("0").rstrip(".") if d != d.to_integral() else str(int(d))
    else:
        require(isinstance(value, str), "Sample ID must be text or numeric.")
        value = value.strip()
    require(bool(value) and len(value) <= 128 and not any(ord(c) < 32 for c in value), "Missing or invalid sample ID.")
    return value


def validate_columns(columns):
    require(isinstance(columns, dict) and set(columns) == set(ALL_FIELDS), "Map all twelve input/result fields explicitly.")
    result = {}
    for field in ALL_FIELDS:
        col = columns[field]
        require(isinstance(col, str) and re.fullmatch(r"[A-Za-z]{1,3}", col) is not None, "Columns must be Excel letters.")
        col = col.upper()
        require(column_index_from_string(col) <= 16384, "Mapped column exceeds Excel limits.")
        result[field] = col
    require(len(set(result.values())) == len(result), "Two fields cannot map to the same column.")
    return result


def _text(cell):
    require(cell.data_type != "f", f"Input formula at {cell.coordinate} requires a values-based input or explicit review.")
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float, Decimal)):
        return normalize_id(value)
    return str(value).strip()


def _amount(cell):
    require(cell.data_type != "f", f"Input formula at {cell.coordinate} requires a values-based input or explicit review.")
    require(cell.value is not None and not isinstance(cell.value, bool), f"Missing amount at {cell.coordinate}.")
    text = str(cell.value).strip()
    # No locale/currency guessing: numeric cells and invariant decimal text only.
    require(re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text) is not None, f"Amount at {cell.coordinate} is not an invariant decimal.")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise WorkbookLayoutError(f"Invalid amount at {cell.coordinate}.") from exc
    require(value.is_finite() and value >= 0, f"Invalid amount at {cell.coordinate}.")
    return format(value, "f")


def _legend_heading(sheet, row, cells, columns):
    """Recognize only the template's merged, otherwise empty legend heading."""
    cell = cells["sample_id"]
    if cell.data_type == "f" or normalize_header(cell.value) != "tickmarklegend":
        return False
    if any(sheet[f"{columns[field]}{row}"].value is not None
           and str(sheet[f"{columns[field]}{row}"].value).strip()
           for field in ALL_FIELDS if field != "sample_id"):
        return False
    return any(merged.min_row == row == merged.max_row
               and merged.min_col == cell.column and merged.max_col > cell.column
               for merged in sheet.merged_cells.ranges)


def _legend_annotation(cells, populated):
    """Accept the known X/dash definitions or an ID-column asterisk footnote.

    This is not a general missing-input exemption. Unexpected data after the
    footer remains an error, including complete records and partial records.
    """
    cell = cells["sample_id"]
    if any(c.data_type == "f" for c in cells.values()):
        return False
    if populated == [True, True, False, False, False]:
        return (str(cell.value).strip(), normalize_header(cells["date_time"].value)) in {
            ("X", "testingattributesatisfiedwithoutexception"),
            ("-", "notapplicable"),
        }
    return (populated == [True, False, False, False, False]
            and isinstance(cell.value, str) and cell.value.lstrip().startswith("*"))


def discover_workbook(path, layout=None):
    """Return a proven sheet/header/column mapping and source-bound sample rows.

    Optional layout: {sheet_name, header_row, columns:{all 12 fields:letter},
    first_data_row?, last_data_row?}. Automatic discovery scans the first 100
    rows / 512 columns of every worksheet and fails on ambiguity.
    """
    path = Path(path).resolve(strict=True)
    require(path.is_file() and path.suffix.lower() == ".xlsx", "Select one .xlsx workbook.")
    before = source_hash(path)
    book = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    warnings = []
    try:
        if layout is not None:
            require(isinstance(layout, dict), "Layout must be an object.")
            name, header = layout.get("sheet_name"), layout.get("header_row")
            require(name in book.sheetnames and type(header) is int and 1 <= header <= 1048576, "Explicit sheet/header mapping is invalid.")
            columns = validate_columns(layout.get("columns"))
            sheet = book[name]
        else:
            aliases = {normalize_header(alias): field for field, names in ALIASES.items() for alias in names}
            candidates, duplicate_candidates = [], []
            for candidate in book.worksheets:
                for row in candidate.iter_rows(min_row=1, max_row=min(candidate.max_row, 100), max_col=min(candidate.max_column, 512)):
                    found, duplicates = {}, set()
                    for cell in row:
                        field = aliases.get(normalize_header(cell.value))
                        if field:
                            if field in found:
                                duplicates.add(field)
                            found[field] = cell.column_letter
                    if set(found) == set(ALL_FIELDS):
                        (duplicate_candidates if duplicates else candidates).append((candidate.title, row[0].row, found))
            require(not duplicate_candidates, "Ambiguous duplicate headers; supply an explicit layout.")
            require(len(candidates) == 1, "No unique complete header row found; supply an explicit sheet/header/column layout.")
            name, header, columns = candidates[0]
            sheet = book[name]
        columns = validate_columns(columns)
        first = (layout or {}).get("first_data_row", header + 1)
        last = (layout or {}).get("last_data_row", sheet.max_row)
        require(type(first) is int and type(last) is int and header < first <= last <= 1048576, "Explicit data row bounds are invalid.")
        samples, seen, legend_row = [], set(), None
        for row in range(first, last + 1):
            cells = {field: sheet[f"{columns[field]}{row}"] for field in INPUT_FIELDS}
            populated = [cells[field].value is not None and str(cells[field].value).strip() != "" for field in INPUT_FIELDS]
            if not any(populated):
                continue
            if legend_row is not None:
                require(_legend_annotation(cells, populated),
                        f"Unexpected input data below the Tickmark Legend at row {row}; supply explicit data row bounds.")
                continue
            if samples and _legend_heading(sheet, row, cells, columns):
                legend_row = row
                warnings.append(f"Recognized Tickmark Legend footer at row {row}; its legend definitions and asterisk footnotes were excluded from sample rows.")
                continue
            if populated[0] and not any(populated[1:]):
                warnings.append(f"Row {row} contains only an ID/label and was not treated as a wire record; use explicit row bounds if needed.")
                continue
            require(populated[0], f"Input data without a sample ID at row {row}.")
            require(cells["sample_id"].data_type != "f", f"Sample ID formula at row {row} is unsupported.")
            sid = normalize_id(cells["sample_id"].value)
            require(sid not in seen, f"Duplicate sample ID at row {row}.")
            seen.add(sid)
            for field in RESULT_FIELDS:
                ref = f"{columns[field]}{row}"
                require(not any(ref in merged for merged in sheet.merged_cells.ranges), f"Result cell {ref} is merged.")
            for field in ("customer", "account", "date_time"):
                require(populated[INPUT_FIELDS.index(field)], f"Missing {field} at row {row}.")
            account = _text(cells["account"])
            if isinstance(cells["account"].value, (int, float)) and len(re.sub(r"\D", "", account)) > 15:
                warnings.append(f"Numeric account at row {row} may already have lost precision in Excel; manual identity review is required.")
            samples.append({"sample_id": sid, "row": row, "date_time": _text(cells["date_time"]),
                            "customer": _text(cells["customer"]), "account": account,
                            "amount_decimal": _amount(cells["amount"])})
        require(bool(samples), "No wire sample rows found below the mapped header.")
        require(source_hash(path) == before, "Workbook changed during discovery.")
        return {"sheet_name": name, "header_row": header, "columns": columns,
                "samples": samples, "source_hash": before, "warnings": warnings}
    finally:
        book.close()
